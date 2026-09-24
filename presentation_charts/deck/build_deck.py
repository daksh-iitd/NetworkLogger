from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image as PILImage

# Palette pulled directly from the Android app's own UI, for consistency
NAVY_BG = RGBColor(0x0A, 0x0E, 0x1A)      # app background
CARD_BG = RGBColor(0x16, 0x1D, 0x2F)      # app card background
INDIGO = RGBColor(0x63, 0x66, 0xF1)       # app accent bar
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x11, 0x18, 0x27)          # near-black body text on light bg
MUTED = RGBColor(0x6B, 0x72, 0x80)        # muted gray, matches app subtitle gray

MAPS = "/Users/dakshgupta/Documents/MTP-1/Android-Telecom-Monitoring-App-main/presentation_charts/maps"
S1 = f"{MAPS}/session1_7b1194d6bf7630cb_20260907_225448"
S2 = f"{MAPS}/session2_c176d0f3ae801a6d_20260902_092536"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def add_slide(bg_color):
    slide = prs.slides.add_slide(BLANK)
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = bg_color
    return slide


def add_text(slide, text, left, top, width, height, size, color, bold=False,
             align=PP_ALIGN.LEFT, font="Calibri", italic=False, anchor=None):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    if anchor:
        tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
        run.font.name = font
    return box


def add_bullets(slide, items, left, top, width, height, size=16, color=INK,
                 heading_color=INDIGO, gap_after=10):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap_after)
        if isinstance(item, tuple):
            heading, rest = item
            r1 = p.add_run()
            r1.text = heading
            r1.font.bold = True
            r1.font.size = Pt(size)
            r1.font.color.rgb = heading_color
            r1.font.name = "Calibri"
            r2 = p.add_run()
            r2.text = rest
            r2.font.size = Pt(size)
            r2.font.color.rgb = color
            r2.font.name = "Calibri"
        else:
            r = p.add_run()
            r.text = item
            r.font.size = Pt(size)
            r.font.color.rgb = color
            r.font.name = "Calibri"
    return box


def add_pagenum(slide, n):
    add_text(slide, str(n), 12.6, 7.05, 0.5, 0.35, 11, MUTED, align=PP_ALIGN.RIGHT)


def add_fitted_picture(slide, path, col_x, col_w, y, max_h):
    im = PILImage.open(path)
    aspect = im.size[0] / im.size[1]
    h = max_h
    w = h * aspect
    if w > col_w:
        w = col_w
        h = w / aspect
    x = col_x + (col_w - w) / 2
    y_centered = y + (max_h - h) / 2
    slide.shapes.add_picture(path, Inches(x), Inches(y_centered), width=Inches(w), height=Inches(h))
    return h


def add_pair_slide(page_no, title, img1, img2, caption1, caption2):
    s = add_slide(WHITE)
    add_text(s, title, 0.7, 0.4, 11.5, 0.7, 30, INK, bold=True)
    col_w = 5.7
    y_img = 1.35
    max_h = 4.5
    x1 = 0.65
    x2 = x1 + col_w + 0.55
    add_fitted_picture(s, img1, x1, col_w, y_img, max_h)
    add_fitted_picture(s, img2, x2, col_w, y_img, max_h)
    add_text(s, caption1, x1, y_img + max_h + 0.08, col_w, 0.6, 13, MUTED, align=PP_ALIGN.CENTER)
    add_text(s, caption2, x2, y_img + max_h + 0.08, col_w, 0.6, 13, MUTED, align=PP_ALIGN.CENTER)
    add_pagenum(s, page_no)


# ============================================================
# Slide 1 - Title
# ============================================================
s = add_slide(NAVY_BG)
add_text(s, "NetworkLogger", 1.0, 2.6, 11.3, 1.0, 44, WHITE, bold=True)
add_text(s, "Android Telecom & Device Health Monitoring", 1.0, 3.5, 11.3, 0.6, 22, RGBColor(0xCA, 0xDC, 0xFC))
add_text(s, "Real-time signal, cell-tower, throughput and device-health logging for network research",
         1.0, 4.15, 10.5, 0.5, 14, MUTED, italic=True)

# ============================================================
# Slide 2 - Introduction & Purpose (NEW)
# ============================================================
s = add_slide(WHITE)
add_text(s, "Introduction & Purpose", 0.7, 0.5, 11.5, 0.8, 32, INK, bold=True)
add_text(s, "A field-data collection tool for real-world cellular network research.",
         0.7, 1.5, 11.7, 0.9, 22, INDIGO, bold=True, italic=True)
add_text(s,
         "Lab equipment and drive tests capture network conditions in controlled settings. "
         "NetworkLogger instead runs as an ordinary Android app on ordinary phones, capturing "
         "authentic signal, mobility, and throughput data during everyday use — the actual "
         "conditions a network operates under.",
         0.7, 2.55, 11.7, 1.3, 16, INK)
add_bullets(s, [
    ("What.  ", "Cellular RF conditions, GPS-based mobility, realized throughput, and device-health context (battery/RAM/storage) that can confound signal readings."),
    ("How.  ", "Two independent logging pipelines on-device — a user-triggered 5-second foreground session, and an always-on 15-minute background trace — both writing structured CSV logs that sync to the cloud."),
    ("Why it matters.  ", "Enables comparing signal quality, radio-technology behavior (4G/5G), and throughput across real devices, carriers, and locations — not synthetic benchmarks."),
], 0.7, 4.0, 11.9, 3.0, size=15, gap_after=14)
add_pagenum(s, 2)

# ============================================================
# Slide 3 - What It Monitors
# ============================================================
s = add_slide(WHITE)
add_text(s, "What It Monitors", 0.7, 0.4, 10, 0.7, 32, INK, bold=True)
add_bullets(s, [
    ("Cellular signal & cell identity  —  ", "RSRP, RSRQ, PCI, TAC, NCI; distinguishes 4G, 5G Standalone, and 5G NSA"),
    ("Location & throughput  —  ", "GPS lat/lon, downlink/uplink kbps"),
    ("Device health  —  ", "battery %, RAM, storage"),
    ("Two logging cadences  —  ", "a 5-second foreground service for fine-grained sessions, and a 15-minute background worker for always-on monitoring"),
], 0.7, 1.25, 11.8, 3.0, size=17, gap_after=13)

glossary_box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(4.55), Inches(11.9), Inches(2.35))
glossary_box.fill.solid()
glossary_box.fill.fore_color.rgb = RGBColor(0xF5, 0xF6, 0xFA)
glossary_box.line.fill.background()
glossary_box.shadow.inherit = False
add_text(s, "QUICK GLOSSARY", 1.05, 4.75, 6, 0.35, 13, INDIGO, bold=True)
add_bullets(s, [
    ("RSRP  —  ", "how strong the signal is"),
    ("RSRQ  —  ", "how clean/usable the signal is (accounts for interference)"),
    ("PCI  —  ", "which physical cell/sector you're connected to"),
], 1.05, 5.15, 5.6, 1.6, size=13, gap_after=8)
add_bullets(s, [
    ("TAC  —  ", "identifies a group of nearby cells"),
    ("NCI  —  ", "the exact cell's unique ID"),
], 6.9, 5.15, 5.5, 1.6, size=13, gap_after=8)
add_pagenum(s, 3)

# ============================================================
# Slide 4 - System Architecture
# ============================================================
s = add_slide(WHITE)
add_text(s, "System Architecture", 0.7, 0.5, 10, 0.8, 32, INK, bold=True)

stages = ["Android App\n(Service + WorkManager)", "Firebase\nCloud Storage", "Local Ingestion\nScript (Python)", "InfluxDB", "Grafana"]
n = len(stages)
box_w, box_h = 2.05, 1.3
gap = 0.35
total_w = n * box_w + (n - 1) * gap
start_x = (13.333 - total_w) / 2
y = 3.0

for i, label in enumerate(stages):
    x = start_x + i * (box_w + gap)
    shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(box_w), Inches(box_h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD_BG
    shape.line.fill.background()
    shape.shadow.inherit = False
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = label
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = WHITE
    if i < n - 1:
        arrow_x = x + box_w
        add_text(s, "→", arrow_x, y, gap, box_h, 22, INDIGO, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

add_bullets(s, [
    ("1.  ", "Tapping Start Logging launches a foreground Service that samples GPS, signal, and throughput every 5 seconds and writes to a local CSV."),
    ("2.  ", "Independently, a WorkManager job runs every 15 minutes in the background — capturing device health and coarse signal data whether or not a session is active."),
    ("3.  ", "On Stop, both CSVs are tagged with device/session/tester metadata, uploaded to Firebase Cloud Storage, and cleared locally for the next session."),
], 0.9, 4.85, 11.5, 2.5, size=14, gap_after=8)
add_pagenum(s, 4)

# ============================================================
# Slide 5 - Engineering Challenges Solved (condensed, one slide)
# ============================================================
s = add_slide(WHITE)
add_text(s, "Engineering Challenges Solved", 0.7, 0.5, 11.5, 0.8, 32, INK, bold=True)
add_bullets(s, [
    ("Background execution limits (Doze).  ",
     "The 5-second sampling loop was silently throttled once the screen turned off — gaps grew to several minutes. Fixed with a dedicated foreground Service plus a WorkManager-based watchdog that detects and alerts on silent failures."),
    ("Correct radio-technology detection.  ",
     "TelephonyManager.networkType misreports 5G NSA (LTE anchor + NR data) as plain 4G. Switched to TelephonyDisplayInfo.overrideNetworkType, the modern API built to distinguish NSA / SA / carrier-aggregation states."),
    ("Correct Cell ID extraction.  ",
     "CellIdentityLte.ci returns the ECI, not the local Cell ID. Decomposed per 3GPP: ECI = 256 × eNodeB ID + Cell ID."),
], 0.7, 1.6, 11.9, 4.8, size=17, gap_after=20)
add_pagenum(s, 5)

# ============================================================
# Slide 6 - Data & Findings: Route Map
# ============================================================
s = add_slide(WHITE)
add_text(s, "Data & Findings — Route", 0.7, 0.4, 11.5, 0.7, 32, INK, bold=True)
add_fitted_picture(s, f"{S1}/1_route_map.png", 3.67, 6.0, 1.3, 5.3)
add_text(s, "Route colored by elapsed time — reveals repeated test passes over ~4 hours, not just a single trip",
         2.0, 6.75, 9.33, 0.5, 14, MUTED, align=PP_ALIGN.CENTER)
add_pagenum(s, 6)

# ============================================================
# Slides 7-9 - Data & Findings: RSRP, Network Type, Throughput (both sessions)
# ============================================================
add_pair_slide(
    7, "Data & Findings — Signal Strength (RSRP)",
    f"{S1}/2_route_rsrp.png", f"{S2}/2_route_rsrp.png",
    "Session 1: ~10km journey, mixed coverage quality",
    "Session 2: short drive, sharp dead zone at low point"
)

add_pair_slide(
    8, "Data & Findings — Network Type (4G / 5G)",
    f"{S1}/3_route_network_type.png", f"{S2}/3_route_network_type.png",
    "Session 1: mixed 4G/5G along the route",
    "Session 2: consistently 4G throughout"
)

add_pair_slide(
    9, "Data & Findings — Downlink Throughput",
    f"{S1}/4_route_throughput.png", f"{S2}/4_route_throughput.png",
    "Session 1: bursty throughput, log scale",
    "Session 2: mostly low-rate background traffic"
)

# ============================================================
# Slide 10 - Live Demo
# ============================================================
s = add_slide(NAVY_BG)
add_text(s, "Live Demo", 1.0, 2.7, 11, 1.0, 40, WHITE, bold=True)
add_text(s, "Grafana + InfluxDB dashboard", 1.0, 3.6, 11, 0.6, 20, RGBColor(0xCA, 0xDC, 0xFC))
add_text(s, "Real-time signal, throughput, and battery data across every device and session",
         1.0, 4.2, 10.5, 0.5, 14, MUTED, italic=True)

out_path = "/Users/dakshgupta/Documents/MTP-1/Android-Telecom-Monitoring-App-main/presentation_charts/NetworkLogger_Presentation.pptx"
prs.save(out_path)
print("Saved:", out_path)
print("Total slides:", len(prs.slides))
